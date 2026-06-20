from django.shortcuts import render, get_object_or_404, redirect
from django.contrib.auth.decorators import login_required, user_passes_test
from django.contrib import messages
from django.utils import timezone
from datetime import timedelta
from django.db.models import Sum, F, Q
from django.http import HttpResponse, JsonResponse
from django.db import models, transaction
import openpyxl
from io import BytesIO
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph
from reportlab.lib.styles import getSampleStyleSheet
from .models import Producto, MovimientoInventario, CategoriaProducto, Empresa
from TratamientoApp.models import Tratamiento, TratamientoProducto
from .forms import ProductoForm, EntradaStockForm, SalidaStockForm, EmpresaForm
from CuentasApp.models import Usuario


def es_auxiliar(user):
    if not user.is_authenticated:
        return False
    return (
        getattr(user, 'id_rol', None)
        and user.id_rol.nombre_rol == 'Auxiliar de Bodega'
    ) or user.is_superuser


def es_admin(user):
    return user.is_authenticated and (
        getattr(user, 'is_admin', False) or user.is_superuser
    )


def es_auxiliar_o_admin(user):
    return es_auxiliar(user) or es_admin(user)


# ─── Dashboard ────────────────────────────────────────────────────────────────

@login_required
@user_passes_test(es_auxiliar)
def dashboard_auxiliar(request):
    usuario_actual = request.user
    productos_bajo_stock = Producto.objects.filter(
        stock_actual__lte=models.F('stock_minimo'), activo=1
    )
    hoy = timezone.now().date()
    proximo_mes = hoy + timedelta(days=30)
    productos_por_vencer = Producto.objects.filter(
        fecha_vencimiento__range=[hoy, proximo_mes], activo=1
    )
    total_productos = Producto.objects.filter(activo=1).count()
    ultimos_movimientos = MovimientoInventario.objects.all().order_by('-fecha_movimiento')[:5]

    return render(request, 'InventarioApp/dashboard_auxiliar.html', {
        'usuario_perfil': usuario_actual,
        'bajo_stock': productos_bajo_stock,
        'por_vencer': productos_por_vencer,
        'total_productos': total_productos,
        'ultimos_movimientos': ultimos_movimientos,
    })


# ─── Lista de Inventario ──────────────────────────────────────────────────────

@login_required
def lista_inventario(request):
    categorias = CategoriaProducto.objects.all()
    # Todas las empresas para el filtro y los modales
    empresas = Empresa.objects.all().order_by('nombre_empresa')

    productos = Producto.objects.all().select_related('id_categoria').order_by('nombre_producto')

    nombre        = request.GET.get('nombre')
    categoria     = request.GET.get('categoria')
    stock_min     = request.GET.get('stock_min')
    stock_max     = request.GET.get('stock_max')
    estado        = request.GET.get('estado')
    empresa_filtro = request.GET.get('empresa')

    if nombre:
        if nombre.startswith('#'):
            cod_limpio = nombre.replace('#', '').strip()
            productos = productos.filter(codigo_producto__icontains=cod_limpio)
        else:
            productos = productos.filter(
                Q(nombre_producto__icontains=nombre) | Q(codigo_producto__icontains=nombre)
            )
    if categoria:
        productos = productos.filter(id_categoria__id_categoria=categoria)
    if stock_min:
        productos = productos.filter(stock_actual__gte=stock_min)
    if stock_max:
        productos = productos.filter(stock_actual__lte=stock_max)
    if estado == 'activo':
        productos = productos.filter(activo=1)
    if estado == 'inactivo':
        productos = productos.filter(activo=0)
    if estado == 'critico':
        productos = productos.filter(stock_actual__lte=F('stock_minimo'))

    # Req 1.1.4: filtrar por empresa a través de MovimientoInventario ENTRADA
    if empresa_filtro:
        ids = MovimientoInventario.objects.filter(
            tipo_movimiento='ENTRADA',
            empresa_asociada__id_empresa=empresa_filtro
        ).values_list('id_producto_id', flat=True).distinct()
        productos = productos.filter(id_producto__in=ids)

    # ── Excel ────────────────────────────────────────────────────────────────
    if request.GET.get('exportar') == 'excel':
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Inventario"
        ws.append(['Código', 'Producto', 'Categoría', 'Stock', 'Mínimo', 'P.Venta', 'Vencimiento', 'Estado'])
        for p in productos:
            ws.append([
                p.codigo_producto or 'S/C', p.nombre_producto,
                p.id_categoria.nombre_categoria, p.stock_actual, p.stock_minimo,
                str(p.precio_venta),
                p.fecha_vencimiento.strftime('%d/%m/%Y') if p.fecha_vencimiento else 'N/A',
                'Activo' if p.activo == 1 else 'Inactivo',
            ])
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename="Inventario.xlsx"'
        wb.save(response)
        return response

    # ── PDF ──────────────────────────────────────────────────────────────────
    if request.GET.get('exportar') == 'pdf':
        buffer = BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=letter)
        elements = []
        styles = getSampleStyleSheet()
        elements.append(Paragraph("Inventario de Productos - OdontoClinick", styles['Title']))
        elements.append(Paragraph(f"Generado: {timezone.now().strftime('%d/%m/%Y %H:%M')}", styles['Normal']))
        data = [['Producto', 'Categoría', 'Stock', 'Mínimo', 'P.Venta', 'Estado']]
        for p in productos:
            data.append([p.nombre_producto[:25], p.id_categoria.nombre_categoria,
                         str(p.stock_actual), str(p.stock_minimo),
                         f"${p.precio_venta}", 'Activo' if p.activo == 1 else 'Inactivo'])
        t = Table(data, colWidths=[140, 90, 50, 50, 70, 60])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#3f2e23')),
            ('TEXTCOLOR', (0,0), (-1,0), colors.whitesmoke),
            ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
            ('GRID', (0,0), (-1,-1), 0.5, colors.grey),
            ('FONTSIZE', (0,0), (-1,-1), 8),
        ]))
        elements.append(t)
        doc.build(elements)
        response = HttpResponse(content_type='application/pdf')
        response['Content-Disposition'] = 'attachment; filename="Inventario.pdf"'
        response.write(buffer.getvalue())
        buffer.close()
        return response

    return render(request, 'InventarioApp/lista_inventario.html', {
        'productos': productos,
        'categorias': categorias,
        'empresas': empresas,
        'total': productos.count(),
    })


# ─── AJAX: buscar usuario por nombre para el modal de salida ──────────────────

@login_required
def buscar_usuario_ajax(request):
    """
    GET /inventario/buscar-usuario/?q=texto
    Devuelve JSON con lista de usuarios que coincidan con nombre o apellidos.
    """
    q = request.GET.get('q', '').strip()
    if len(q) < 2:
        return JsonResponse({'usuarios': []})
    usuarios = Usuario.objects.filter(
        Q(nombre__icontains=q) | Q(apellidos__icontains=q) | Q(nombre_usuario__icontains=q)
    ).values('id_usuario', 'nombre', 'apellidos', 'nombre_usuario')[:10]
    return JsonResponse({'usuarios': list(usuarios)})


# ─── CRUD Productos ───────────────────────────────────────────────────────────

@login_required
def crear_producto(request):
    if request.method == 'POST':
        form = ProductoForm(request.POST)
        if form.is_valid():
            try:
                producto = form.save(commit=False)
                producto.activo = 1
                producto.save()
                messages.success(request, f"Producto '{producto.nombre_producto}' creado con éxito.")
                return redirect('lista_inventario')
            except Exception as e:
                messages.error(request, f"Error inesperado: {e}")
        else:
            messages.error(request, "Por favor, corrige los errores en el formulario.")
    else:
        form = ProductoForm()
    return render(request, 'InventarioApp/crear_producto.html', {'form': form})


@login_required
def editar_producto(request, id_producto):
    producto = get_object_or_404(Producto, id_producto=id_producto)
    if request.method == 'POST':
        form = ProductoForm(request.POST, instance=producto)
        if form.is_valid():
            form.save()
            messages.success(request, f"¡Producto '{producto.nombre_producto}' actualizado!")
            return redirect('lista_inventario')
    else:
        form = ProductoForm(instance=producto)
    return render(request, 'InventarioApp/crear_producto.html', {
        'form': form, 'editando': True, 'producto': producto
    })


@login_required
def alternar_estado_producto(request, producto_id):
    if request.method == 'POST':
        try:
            producto = Producto.objects.get(id_producto=producto_id)
            producto.activo = 0 if producto.activo == 1 else 1
            producto.save()
            return JsonResponse({'status': 'ok', 'nuevo_estado': producto.activo})
        except Producto.DoesNotExist:
            return JsonResponse({'status': 'error'}, status=404)
    return JsonResponse({'status': 'error'}, status=400)


# ─── CRUD Empresas (Req nuevo: auxiliar + admin) ──────────────────────────────

@login_required
@user_passes_test(es_auxiliar_o_admin)
def lista_empresas(request):
    empresas = Empresa.objects.all().order_by('nombre_empresa')

    # Filtros
    nombre_q    = request.GET.get('nombre', '').strip()
    tipo_filtro = request.GET.get('tipo', '')

    if nombre_q:
        empresas = empresas.filter(
            Q(nombre_empresa__icontains=nombre_q) | Q(nit__icontains=nombre_q)
        )
    if tipo_filtro == 'proveedor':
        empresas = empresas.filter(es_proveedor=True)
    elif tipo_filtro == 'comprador':
        empresas = empresas.filter(es_comprador=True)

    # Excel
    if request.GET.get('exportar') == 'excel':
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Empresas"
        ws.append(['ID', 'Nombre', 'NIT', 'Proveedor', 'Comprador'])
        for e in empresas:
            ws.append([e.id_empresa, e.nombre_empresa, e.nit,
                       'Sí' if e.es_proveedor else 'No',
                       'Sí' if e.es_comprador else 'No'])
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename="Empresas.xlsx"'
        wb.save(response)
        return response

    return render(request, 'InventarioApp/lista_empresas.html', {
        'empresas': empresas,
        'total': empresas.count(),
    })


@login_required
@user_passes_test(es_auxiliar_o_admin)
def crear_empresa(request):
    """Creación vía modal (POST AJAX o form normal)."""
    if request.method == 'POST':
        form = EmpresaForm(request.POST)
        if form.is_valid():
            empresa = form.save()
            if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                return JsonResponse({
                    'status': 'ok',
                    'id': empresa.id_empresa,
                    'nombre': empresa.nombre_empresa,
                })
            messages.success(request, f"Empresa '{empresa.nombre_empresa}' creada.")
            return redirect('lista_empresas')
        else:
            if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                return JsonResponse({'status': 'error', 'errors': form.errors}, status=400)
            messages.error(request, "Corrige los errores del formulario.")
    else:
        form = EmpresaForm()
    return render(request, 'InventarioApp/lista_empresas.html', {'form': form})


@login_required
@user_passes_test(es_auxiliar_o_admin)
def editar_empresa(request, pk):
    empresa = get_object_or_404(Empresa, pk=pk)
    if request.method == 'POST':
        form = EmpresaForm(request.POST, instance=empresa)
        if form.is_valid():
            form.save()
            if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                return JsonResponse({'status': 'ok', 'nombre': empresa.nombre_empresa})
            messages.success(request, f"Empresa '{empresa.nombre_empresa}' actualizada.")
            return redirect('lista_empresas')
        else:
            if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                return JsonResponse({'status': 'error', 'errors': form.errors}, status=400)
    else:
        # GET: devolver datos para el modal de edición
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return JsonResponse({
                'id': empresa.id_empresa,
                'nombre_empresa': empresa.nombre_empresa,
                'nit': empresa.nit,
                'es_proveedor': empresa.es_proveedor,
                'es_comprador': empresa.es_comprador,
            })
    return redirect('lista_empresas')


@login_required
@user_passes_test(es_auxiliar_o_admin)
def alternar_tipo_empresa(request, pk):
    """Alterna es_proveedor o es_comprador vía AJAX."""
    if request.method == 'POST':
        import json
        empresa = get_object_or_404(Empresa, pk=pk)
        data = json.loads(request.body)
        campo = data.get('campo')  # 'es_proveedor' o 'es_comprador'
        if campo in ('es_proveedor', 'es_comprador'):
            setattr(empresa, campo, not getattr(empresa, campo))
            empresa.save()
            return JsonResponse({'status': 'ok', 'valor': getattr(empresa, campo)})
    return JsonResponse({'status': 'error'}, status=400)


# ─── Tratamientos ─────────────────────────────────────────────────────────────

@login_required
def lista_tratamientos_auxiliar(request):
    txt_buscar = request.GET.get('buscar', '').strip()
    tratamientos = Tratamiento.objects.filter(activo=1)
    if txt_buscar:
        tratamientos = tratamientos.filter(nombre_tratamiento__icontains=txt_buscar)
    tratamientos = tratamientos.order_by('nombre_tratamiento')
    return render(request, 'InventarioApp/lista_tratamientos_auxiliar.html', {'tratamientos': tratamientos})


@login_required
def gestionar_insumos(request, pk):
    tratamiento = get_object_or_404(Tratamiento, pk=pk)
    if request.method == 'POST':
        producto_id = request.POST.get('producto')
        cantidad    = request.POST.get('cantidad')
        producto    = get_object_or_404(Producto, id_producto=producto_id)
        TratamientoProducto.objects.update_or_create(
            id_tratamiento=tratamiento,
            id_producto=producto,
            defaults={'cantidad_requerida': cantidad}
        )
        messages.success(request, f"Insumo vinculado a {tratamiento.nombre_tratamiento}")
        return redirect('gestionar_insumos', pk=pk)

    insumos  = TratamientoProducto.objects.filter(id_tratamiento=tratamiento).select_related('id_producto')
    productos = Producto.objects.filter(activo=1)
    return render(request, 'InventarioApp/gestionar_insumos.html', {
        'tratamiento': tratamiento,
        'insumos': insumos,
        'productos': productos,
    })


@login_required
def eliminar_insumo(request, pk):
    relacion = get_object_or_404(TratamientoProducto, pk=pk)
    id_trat  = relacion.id_tratamiento.pk
    relacion.delete()
    messages.error(request, "Insumo eliminado del tratamiento.")
    return redirect('gestionar_insumos', pk=id_trat)


# ─── Movimientos de Bodega ────────────────────────────────────────────────────

@login_required
def entrada_stock(request, pk):
    producto_obj = get_object_or_404(Producto, id_producto=pk)
    if request.method == 'POST':
        form = EntradaStockForm(request.POST)
        if form.is_valid():
            try:
                with transaction.atomic():
                    cantidad = form.cleaned_data['cantidad']
                    motivo   = form.cleaned_data.get('motivo') or 'Entrada de stock'
                    empresa  = form.cleaned_data.get('empresa')
                    precio   = form.cleaned_data.get('precio_transaccion')

                    stock_anterior = producto_obj.stock_actual
                    producto_obj.stock_actual += cantidad
                    producto_obj.save()

                    MovimientoInventario.objects.create(
                        id_producto=producto_obj,
                        tipo_movimiento='ENTRADA',
                        cantidad=cantidad,
                        stock_anterior=stock_anterior,
                        stock_nuevo=producto_obj.stock_actual,
                        motivo=motivo,
                        id_usuario=request.user,
                        empresa_asociada=empresa,
                        precio_transaccion=precio,
                    )

                    if empresa and precio is not None:
                        from .models import ProductoEmpresa
                        pe, created = ProductoEmpresa.objects.get_or_create(
                            id_producto=producto_obj,
                            id_empresa=empresa,
                            precio_compra_proveedor=precio,
                            defaults={'stock_proveedor': cantidad},
                        )
                        if not created:
                            pe.stock_proveedor += cantidad
                            pe.save()

                messages.success(
                    request,
                    f"✓ Entrada de {cantidad} uds. registrada"
                    + (f" — {empresa.nombre_empresa} @ ${precio}/ud." if empresa else ".")
                )
            except Exception as e:
                messages.error(request, f"Error al procesar la entrada: {e}")
        else:
            for field, errs in form.errors.items():
                for err in errs:
                    messages.error(request, err)
    return redirect('lista_inventario')


@login_required
def salida_stock(request, pk):
    producto_obj = get_object_or_404(Producto, id_producto=pk)
    if request.method == 'POST':
        form = SalidaStockForm(request.POST)
        if form.is_valid():
            cantidad   = form.cleaned_data['cantidad']
            motivo     = form.cleaned_data.get('motivo') or 'Salida de stock'
            empresa    = form.cleaned_data.get('empresa')
            cliente_id = form.cleaned_data.get('cliente_id')
            precio     = form.cleaned_data.get('precio_transaccion')

            if cantidad > producto_obj.stock_actual:
                messages.error(request, f"Stock insuficiente. Disponible: {producto_obj.stock_actual} uds.")
                return redirect('lista_inventario')

            try:
                with transaction.atomic():
                    cliente_obj = None
                    if cliente_id:
                        try:
                            cliente_obj = Usuario.objects.get(pk=cliente_id)
                        except Usuario.DoesNotExist:
                            messages.error(request, f"No existe el usuario con ID {cliente_id}.")
                            return redirect('lista_inventario')

                    stock_anterior = producto_obj.stock_actual
                    producto_obj.stock_actual -= cantidad
                    producto_obj.save()

                    MovimientoInventario.objects.create(
                        id_producto=producto_obj,
                        tipo_movimiento='SALIDA',
                        cantidad=cantidad,
                        stock_anterior=stock_anterior,
                        stock_nuevo=producto_obj.stock_actual,
                        motivo=motivo,
                        id_usuario=request.user,
                        empresa_asociada=empresa,
                        cliente_externo=cliente_obj,
                        precio_transaccion=precio,
                    )

                destino = f" → {empresa.nombre_empresa}" if empresa else (f" → {cliente_obj}" if cliente_obj else "")
                messages.warning(request,
                    f"Salida de {cantidad} uds. registrada{destino}"
                    + (f" @ ${precio}/ud." if precio else "."))
            except Exception as e:
                messages.error(request, f"Error al procesar la salida: {e}")
        else:
            for field, errs in form.errors.items():
                for err in errs:
                    messages.error(request, err)
    return redirect('lista_inventario')


# ─── Informes ─────────────────────────────────────────────────────────────────

@login_required
def informes_avanzados(request):
    productos = Producto.objects.filter(activo=1)
    valor_total = productos.aggregate(total=Sum(F('precio_venta') * F('stock_actual')))['total'] or 0
    primer_dia  = timezone.now().replace(day=1, hour=0, minute=0, second=0, microsecond=0)
    inversion_mes = MovimientoInventario.objects.filter(
        tipo_movimiento='ENTRADA', fecha_movimiento__gte=primer_dia
    ).aggregate(total=Sum(F('cantidad') * F('id_producto__precio_venta')))['total'] or 0

    hoy = timezone.now()
    labels, datos = [], []
    for i in range(6, -1, -1):
        fecha = hoy - timedelta(days=i)
        labels.append(fecha.strftime('%d/%m'))
        total = MovimientoInventario.objects.filter(
            tipo_movimiento='SALIDA', fecha_movimiento__date=fecha.date()
        ).aggregate(Sum('cantidad'))['cantidad__sum'] or 0
        datos.append(total)

    return render(request, 'InventarioApp/informes.html', {
        'valor_total': valor_total,
        'inversion_mes': inversion_mes,
        'labels_grafica': labels,
        'datos_grafica': datos,
        'bajo_stock_count': Producto.objects.filter(stock_actual__lte=F('stock_minimo'), activo=1).count(),
        'insumos_criticos': Producto.objects.filter(stock_actual__lte=F('stock_minimo'), activo=1),
    })


# ─── Kardex ───────────────────────────────────────────────────────────────────

@login_required
def historial_kardex(request):
    movimientos = MovimientoInventario.objects.all().select_related(
        'id_producto', 'id_usuario', 'empresa_asociada', 'cliente_externo'
    ).order_by('-fecha_movimiento')

    producto_id          = request.GET.get('producto_id')
    producto_seleccionado = None
    if producto_id:
        producto_seleccionado = get_object_or_404(Producto, id_producto=producto_id)
        movimientos = movimientos.filter(id_producto=producto_seleccionado)

    fecha_inicio  = request.GET.get('fecha_inicio')
    fecha_fin     = request.GET.get('fecha_fin')
    tipo          = request.GET.get('tipo')
    usuario_q     = request.GET.get('usuario')
    query         = request.GET.get('producto', '').strip()
    empresa_filtro = request.GET.get('empresa')

    if query and not producto_id:
        if query.startswith('#'):
            movimientos = movimientos.filter(id_producto__codigo_producto__icontains=query.replace('#',''))
        else:
            movimientos = movimientos.filter(id_producto__nombre_producto__icontains=query)
    if fecha_inicio:
        movimientos = movimientos.filter(fecha_movimiento__date__gte=fecha_inicio)
    if fecha_fin:
        movimientos = movimientos.filter(fecha_movimiento__date__lte=fecha_fin)
    if tipo:
        movimientos = movimientos.filter(tipo_movimiento=tipo)
    if usuario_q:
        movimientos = movimientos.filter(
            Q(id_usuario__nombre__icontains=usuario_q) |
            Q(id_usuario__apellidos__icontains=usuario_q)
        )
    if empresa_filtro:
        movimientos = movimientos.filter(empresa_asociada__id_empresa=empresa_filtro)

    # Excel
    if request.GET.get('exportar') == 'excel':
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Kardex"
        ws.append(['Fecha','Producto','Tipo','Cantidad','Stock Ant.','Stock Nuevo','Empresa/Cliente','Precio','Motivo','Responsable'])
        for m in movimientos:
            destino = m.empresa_asociada.nombre_empresa if m.empresa_asociada else (str(m.cliente_externo) if m.cliente_externo else '—')
            ws.append([
                m.fecha_movimiento.strftime('%d/%m/%Y %H:%M') if m.fecha_movimiento else 'N/A',
                m.id_producto.nombre_producto if m.id_producto else 'N/A',
                m.tipo_movimiento, m.cantidad, m.stock_anterior, m.stock_nuevo,
                destino, f"${m.precio_transaccion}" if m.precio_transaccion else '—',
                m.motivo or 'Sin motivo', str(m.id_usuario) if m.id_usuario else 'Sistema',
            ])
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename="Kardex.xlsx"'
        wb.save(response)
        return response

    # PDF
    if request.GET.get('exportar') == 'pdf':
        buffer = BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=letter)
        elements = []
        styles = getSampleStyleSheet()
        titulo = f"Historial de {producto_seleccionado.nombre_producto}" if producto_seleccionado else "Historial de Movimientos"
        elements.append(Paragraph(titulo, styles['Title']))
        elements.append(Paragraph(f"Generado: {timezone.now().strftime('%d/%m/%Y %H:%M')}", styles['Normal']))
        data = [['Fecha','Producto','Tipo','Cant.','Stock Nuevo','Empresa/Cliente','Responsable']]
        for m in movimientos:
            destino = m.empresa_asociada.nombre_empresa[:15] if m.empresa_asociada else (str(m.cliente_externo)[:15] if m.cliente_externo else '—')
            data.append([
                m.fecha_movimiento.strftime('%d/%m/%y') if m.fecha_movimiento else 'N/A',
                m.id_producto.nombre_producto[:18] if m.id_producto else 'N/A',
                m.tipo_movimiento, str(m.cantidad), str(m.stock_nuevo), destino,
                str(m.id_usuario) if m.id_usuario else 'Sistema',
            ])
        t = Table(data, colWidths=[60,110,55,35,55,80,80])
        t.setStyle(TableStyle([
            ('BACKGROUND',(0,0),(-1,0),colors.HexColor('#3f2e23')),
            ('TEXTCOLOR',(0,0),(-1,0),colors.whitesmoke),
            ('FONTNAME',(0,0),(-1,0),'Helvetica-Bold'),
            ('ALIGN',(0,0),(-1,-1),'CENTER'),
            ('GRID',(0,0),(-1,-1),0.5,colors.grey),
            ('FONTSIZE',(0,0),(-1,-1),7),
        ]))
        elements.append(t)
        doc.build(elements)
        response = HttpResponse(content_type='application/pdf')
        response['Content-Disposition'] = 'attachment; filename="Kardex.pdf"'
        response.write(buffer.getvalue())
        buffer.close()
        return response

    empresas_kardex = Empresa.objects.all().order_by('nombre_empresa')
    return render(request, 'InventarioApp/kardex.html', {
        'movimientos': movimientos,
        'total': movimientos.count(),
        'producto_seleccionado': producto_seleccionado,
        'empresas': empresas_kardex,
    })